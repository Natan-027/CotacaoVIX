"""
Robô de cotação CEASA-ES (Grande Vitória)

O que faz:
  1. Abre o formulário do boletim
  2. Escolhe o mercado "CEASA GRANDE VITÓRIA"
  3. Escolhe a data mais recente disponível
  4. Clica em Ok e lê a tabela de resultados
  5. Acrescenta os dados em cotacoes_ceasa.xlsx (sem duplicar o mesmo dia)

Instalação (uma vez só):
  pip install playwright pandas openpyxl lxml
  playwright install chromium

Uso:
  python cotacao_ceasa.py            # roda normalmente
  python cotacao_ceasa.py --debug    # abre o navegador visível e salva prints/HTML na pasta debug/
"""

import shutil
import sys
from datetime import datetime
from io import StringIO
from pathlib import Path

import pandas as pd
from lxml import html as lxml_html
from openpyxl import Workbook, load_workbook
from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout

URL = "http://200.198.51.71/detec/filtro_boletim_es/filtro_boletim_es.php"
MERCADO = "CEASA GRANDE VITÓRIA"
ARQUIVO = Path("cotacoes_ceasa.xlsx")
NOME_ABA = "Cotações"
CABECALHO = ["Dia", "Mês", "Ano", "Produto", "M.C."]

# Opcional: para gravar só alguns produtos, liste os nomes EXATAMENTE como no site.
# Lista vazia = grava todos os produtos do boletim.
# Exemplo: PRODUTOS = ["AGRIAO", "ALFACE CRESPA", "COUVE", "RUCULA"]
PRODUTOS = []
DEBUG = "--debug" in sys.argv
PASTA_DEBUG = Path("debug")


def salvar_debug(page, nome):
    """Guarda print e HTML da página para facilitar ajustes."""
    PASTA_DEBUG.mkdir(exist_ok=True)
    page.screenshot(path=str(PASTA_DEBUG / f"{nome}.png"), full_page=True)
    (PASTA_DEBUG / f"{nome}.html").write_text(page.content(), encoding="utf-8")


def data_mais_recente(opcoes):
    """opcoes: lista de dicts {value, text}. Devolve a opção com a data mais recente."""
    validas = []
    for op in opcoes:
        texto = op["text"].strip()
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                validas.append((datetime.strptime(texto, fmt), op))
                break
            except ValueError:
                continue
    if not validas:
        raise RuntimeError(f"Nenhuma data reconhecida nas opções: {opcoes}")
    return max(validas, key=lambda x: x[0])[1]


def _um_numero(valor):
    """Texto brasileiro ('1.234,50' ou '4,55') vira número; vazio vira NaN."""
    txt = str(valor).strip().replace(".", "").replace(",", ".")
    try:
        return float(txt)
    except ValueError:
        return float("nan")


NOMES = {"Produtos": "Produto", "MIN": "Mínimo", "MAX": "Máximo"}
COLUNAS_PRECO = ["Mínimo", "M.C.", "Máximo"]


def _texto(el):
    return " ".join(el.text_content().split())


def extrair_produtos(html_str):
    """
    Lê a tabela de produtos do boletim direto do HTML.
    Pega a tabela MAIS INTERNA que tem o cabeçalho 'Produtos ... MIN', ignorando
    as tabelas de layout que envolvem a página. Linhas com um único texto são
    tratadas como título de grupo (se existirem).
    """
    doc = lxml_html.fromstring(html_str)
    linhas_cab = doc.xpath(
        "//tr[*[normalize-space()='Produtos'] and *[normalize-space()='MIN']]"
    )
    tabelas = []
    for tr in linhas_cab:
        t = tr.xpath("ancestor::table[1]")[0]
        if t not in tabelas:
            tabelas.append(t)

    registros, grupos, ignoradas = [], [], 0
    for t in tabelas:
        cabecalho, grupo = None, ""
        for tr in t.xpath("./tr | ./thead/tr | ./tbody/tr | ./tfoot/tr"):
            cels = [_texto(c) for c in tr.xpath("./td | ./th")]
            preenchidas = [c for c in cels if c]
            if not preenchidas:
                continue
            if cels[0] == "Produtos":
                cabecalho = [NOMES.get(c, c) for c in cels]
                continue
            if cabecalho is None:
                continue
            if len(preenchidas) == 1:  # linha de título de grupo
                grupo = preenchidas[0]
                if grupo not in grupos:
                    grupos.append(grupo)
                continue
            if len(cels) != len(cabecalho):
                ignoradas += 1  # linha fora do padrão; ignora
                continue
            reg = dict(zip(cabecalho, cels))
            reg["Grupo"] = grupo
            registros.append(reg)

    if ignoradas:
        print(f"Aviso: {ignoradas} linha(s) fora do padrão foram ignoradas.")
    if not registros:
        return None, grupos
    df = pd.DataFrame(registros)
    for col in COLUNAS_PRECO:
        if col in df.columns:
            df[col] = df[col].map(_um_numero)
    df = df.replace("", pd.NA)
    if df["Grupo"].eq("").all() or df["Grupo"].isna().all():
        df = df.drop(columns="Grupo")
    else:
        ordem = ["Grupo"] + [c for c in df.columns if c != "Grupo"]
        df = df[ordem]
    return df, grupos


def achar_tabela(page):
    """Tenta a página e depois os iframes, até achar a tabela de produtos."""
    for alvo in [page] + list(page.frames):
        try:
            df, grupos = extrair_produtos(alvo.content())
        except Exception:
            continue
        if df is not None and len(df) >= 5:
            if grupos:
                print("Grupos encontrados na página:", grupos)
            return df
    return None


def _inteiro(valor):
    try:
        return int(float(str(valor).strip()))
    except (ValueError, TypeError):
        return None


def _ultima_linha(ws):
    """Última linha que realmente tem conteúdo (ignora linhas vazias só formatadas)."""
    for r in range(ws.max_row, 0, -1):
        if any(c.value not in (None, "") for c in ws[r]):
            return r
    return 0


def _escrever_linha(ws, r, valores):
    for col, v in enumerate(valores, start=1):
        ws.cell(row=r, column=col, value=v)
    ws.cell(row=r, column=1).number_format = "00"    # dia  -> aparece 09
    ws.cell(row=r, column=2).number_format = "00"    # mês  -> aparece 09
    ws.cell(row=r, column=3).number_format = "0"     # ano
    ws.cell(row=r, column=5).number_format = "0.00"  # preço M.C.


def _nova_planilha(linhas):
    wb = Workbook()
    ws = wb.active
    ws.title = NOME_ABA
    ws.append(CABECALHO)
    for i, linha in enumerate(linhas, start=2):
        _escrever_linha(ws, i, linha)
    for letra, larg in zip("ABCDE", (6, 6, 8, 34, 10)):
        ws.column_dimensions[letra].width = larg
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{max(ws.max_row, 1)}"
    wb.save(ARQUIVO)


def _migrar_formato_antigo():
    """
    Se a planilha existente está em um formato de versão anterior do robô
    (coluna 'Data' única), converte para o formato novo e guarda uma cópia.
    """
    if not ARQUIVO.exists():
        return
    try:
        antigo = pd.read_excel(ARQUIVO)
    except Exception:
        return
    if "Data" not in antigo.columns:
        return  # já está no formato novo

    carimbo = f"{datetime.now():%Y%m%d_%H%M%S}"
    copia = ARQUIVO.with_name(f"cotacoes_ceasa_antigo_{carimbo}.xlsx")
    shutil.copy(ARQUIVO, copia)
    print(f"Cópia da planilha antiga salva em {copia.name}")

    if "Produto" in antigo.columns and "M.C." in antigo.columns:
        datas = pd.to_datetime(antigo["Data"], dayfirst=True, errors="coerce")
        linhas = [
            [int(d.day), int(d.month), int(d.year), p, float(mc)]
            for d, p, mc in zip(datas, antigo["Produto"], antigo["M.C."])
            if pd.notna(d) and pd.notna(p) and pd.notna(mc)
        ]
        _nova_planilha(linhas)
        print(f"Planilha antiga convertida para o formato novo ({len(linhas)} linhas mantidas).")
    else:
        ARQUIVO.unlink()  # formato antigo inutilizável; começa do zero


def salvar_planilha(tabela, data_txt):
    """
    Acrescenta o dia no FINAL da planilha (abaixo da última linha preenchida),
    sem repetir o cabeçalho e sem alterar o que já existe.
    Colunas: Dia | Mês | Ano | Produto | M.C.
    """
    d = datetime.strptime(data_txt, "%d/%m/%Y")

    base = tabela[["Produto", "M.C."]].dropna(subset=["M.C."])
    if PRODUTOS:
        base = base[base["Produto"].isin(PRODUTOS)]
        faltando = [p for p in PRODUTOS if p not in set(base["Produto"])]
        if faltando:
            print("Aviso: não encontrei no boletim de hoje:", faltando)
    novas = [[d.day, d.month, d.year, p, float(mc)]
             for p, mc in zip(base["Produto"], base["M.C."])]
    if not novas:
        print("Nenhum produto para gravar hoje.")
        return False

    try:
        _migrar_formato_antigo()

        if not ARQUIVO.exists():
            _nova_planilha(novas)
            print(f"Planilha criada com {len(novas)} linhas em {ARQUIVO}")
            return True

        wb = load_workbook(ARQUIVO)
        ws = wb[NOME_ABA] if NOME_ABA in wb.sheetnames else wb.active
        ultima = _ultima_linha(ws)

        # já gravou esse dia? (olha as colunas Dia/Mês/Ano)
        for dia, mes, ano in ws.iter_rows(min_row=1, max_row=ultima, max_col=3,
                                          values_only=True):
            if (_inteiro(dia), _inteiro(mes), _inteiro(ano)) == (d.day, d.month, d.year):
                print(f"A data {data_txt} já está na planilha. Nada a fazer.")
                return False

        if ultima == 0:  # planilha vazia: escreve o cabeçalho uma única vez
            for col, nome in enumerate(CABECALHO, start=1):
                ws.cell(row=1, column=col, value=nome)
            ultima = 1

        primeira_nova = ultima + 1
        for i, linha in enumerate(novas, start=primeira_nova):
            _escrever_linha(ws, i, linha)
        if ws.auto_filter.ref:
            ws.auto_filter.ref = f"A1:E{primeira_nova + len(novas) - 1}"
        wb.save(ARQUIVO)
        print(f"{len(novas)} linhas gravadas nas linhas {primeira_nova} a "
              f"{primeira_nova + len(novas) - 1} de {ARQUIVO}")
        return True
    except PermissionError:
        print(f"Não consegui salvar. Feche a planilha {ARQUIVO.name} no Excel e rode de novo.")
        sys.exit(1)


def main():
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=not DEBUG)
        page = browser.new_page()
        page.goto(URL, wait_until="networkidle")

        sel_mercado = page.locator("#id_sc_field_mercado")
        sel_data = page.locator("#id_sc_field_datas")

        # 1) Mercado
        sel_mercado.select_option(label=MERCADO)
        page.wait_for_timeout(2000)  # o site carrega as datas via AJAX

        # 2) Data: pega a mais recente (seleciona pelo valor interno da opção)
        opcoes = sel_data.locator("option").evaluate_all(
            "els => els.map(e => ({value: e.value, text: e.textContent}))"
        )
        opcoes = [o for o in opcoes if o["value"].strip()]
        escolhida = data_mais_recente(opcoes)
        data_alvo = escolhida["text"].strip()
        sel_data.select_option(value=escolhida["value"])
        print(f"Data selecionada: {data_alvo}")

        if DEBUG:
            salvar_debug(page, "1_formulario_preenchido")

        # 3) Botão Ok (é um link com id sub_form_t; o site envia o formulário)
        contexto = page.context
        try:
            with page.expect_navigation(wait_until="load", timeout=30000):
                page.click("#sub_form_t")
        except PWTimeout:
            print("Sem navegação na mesma aba; verificando se abriu outra aba...")
        page.wait_for_timeout(3000)
        page = contexto.pages[-1]  # caso o resultado tenha aberto em nova aba
        page.wait_for_load_state("load")

        if DEBUG:
            salvar_debug(page, "2_resultado")

        # 4) Lê a tabela de produtos (com o navegador ainda aberto)
        tabela = achar_tabela(page)
        browser.close()

    if tabela is None:
        print("Não encontrei a tabela de produtos. Rode com --debug e envie os arquivos da pasta debug/.")
        sys.exit(1)

    # 5) Salva sem duplicar o dia
    salvar_planilha(tabela, data_alvo)


if __name__ == "__main__":
    main()
