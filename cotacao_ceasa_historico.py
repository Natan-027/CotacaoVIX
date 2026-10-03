"""
Robô HISTÓRICO - CEASA-ES Grande Vitória

Percorre TODAS as datas que o site oferece, da mais recente para a mais antiga,
e grava cada dia no fim da planilha cotacoes_ceasa.xlsx (Dia | Mês | Ano | Produto | M.C.).

  * Não precisa adivinhar fim de semana ou feriado: usa a lista de datas do próprio site.
  * Pula as datas que já estão na planilha, então pode ser interrompido (Ctrl+C) e
    rodado de novo quantas vezes quiser: ele continua de onde parou e também
    preenche eventuais buracos.
  * Salva a planilha a cada poucos dias, para não perder o progresso.

Precisa estar na MESMA pasta que o cotacao_ceasa.py (usa as funções dele).

Uso:
  python cotacao_ceasa_historico.py --limite 3 --debug   # teste: só 3 dias, navegador visível
  python cotacao_ceasa_historico.py                      # histórico completo
  python cotacao_ceasa_historico.py --limite 100         # só os 100 dias mais recentes que faltam

Feche a planilha no Excel antes de rodar.
"""

import sys
import time
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout

import cotacao_ceasa as cot

SEM_DADOS = Path("historico_datas_sem_dados.txt")  # datas que o site não tem boletim
SALVAR_A_CADA = 40   # dias gravados entre um salvamento e outro
PAUSA = 1.0          # segundos entre consultas (educado com o servidor)
TENTATIVAS = 3       # tentativas por data antes de desistir dela
MAX_PROBLEMAS_SEGUIDOS = 5  # para tudo se o site falhar tantas datas seguidas
DEBUG = "--debug" in sys.argv
ESPERA_TABELA = 1   # segundos extras esperando a tabela aparecer (o robô de publicação aumenta)
AO_FALHAR = None    # função opcional chamada com a página quando a tabela não é encontrada


def argumento(nome):
    if nome in sys.argv:
        return sys.argv[sys.argv.index(nome) + 1]
    return None


# ---------------------------------------------------------------- site ----
def abrir_formulario(page):
    resposta = page.goto(cot.URL, wait_until="networkidle")
    if resposta is not None and resposta.status >= 400:
        raise RuntimeError(f"site respondeu HTTP {resposta.status} ao abrir o formulário")
    page.locator("#id_sc_field_mercado").select_option(label=cot.MERCADO)


def listar_datas(page):
    """Devolve [(datetime, valor_da_opcao)] com todas as datas do site, da mais recente para a mais antiga."""
    abrir_formulario(page)
    page.wait_for_function(
        "document.querySelectorAll('#id_sc_field_datas option').length > 1", timeout=20000)
    opcoes = page.locator("#id_sc_field_datas option").evaluate_all(
        "els => els.map(e => ({value: e.value, text: e.textContent}))")
    datas = []
    for o in opcoes:
        try:
            datas.append((datetime.strptime(o["text"].strip(), "%d/%m/%Y"), o["value"]))
        except ValueError:
            continue  # "Selecione um item"
    return sorted(datas, key=lambda x: x[0], reverse=True)


def consultar_dia(page, valor):
    """Abre o formulário, escolhe a data, clica Ok e devolve a tabela de produtos (ou None)."""
    abrir_formulario(page)
    page.wait_for_function(
        "(v) => Array.from(document.querySelectorAll('#id_sc_field_datas option')).some(o => o.value === v)",
        arg=valor, timeout=20000)
    page.locator("#id_sc_field_datas").select_option(value=valor)
    resposta = None
    try:
        with page.expect_navigation(wait_until="load", timeout=30000) as navegacao:
            page.click("#sub_form_t")
        resposta = navegacao.value
    except PWTimeout:
        pass  # talvez tenha aberto em outra aba; conferimos abaixo
    if resposta is not None and resposta.status >= 400:
        raise RuntimeError(f"site respondeu HTTP {resposta.status} ao consultar a data")
    page.wait_for_timeout(1500)
    pagina = page.context.pages[-1]
    pagina.wait_for_load_state("load")
    try:
        # espera a tabela aparecer (a página pode montá-la depois de carregar)
        tabela = cot.achar_tabela(pagina)
        for _ in range(ESPERA_TABELA):
            if tabela is not None:
                break
            pagina.wait_for_timeout(1000)
            tabela = cot.achar_tabela(pagina)
        if tabela is None and AO_FALHAR is not None:
            AO_FALHAR(pagina)  # gancho de diagnóstico (usado pelo robô de publicação)
        return tabela
    finally:
        if pagina is not page:
            pagina.close()


# ------------------------------------------------------------ planilha ----
def abrir_planilha():
    """Abre a planilha (criando/convertendo se preciso) e devolve wb, ws, última linha e datas já gravadas."""
    cot._migrar_formato_antigo()
    if not cot.ARQUIVO.exists():
        cot._nova_planilha([])
    wb = load_workbook(cot.ARQUIVO)
    ws = wb[cot.NOME_ABA] if cot.NOME_ABA in wb.sheetnames else wb.active
    ultima = cot._ultima_linha(ws)
    if ultima == 0:
        for col, nome in enumerate(cot.CABECALHO, start=1):
            ws.cell(row=1, column=col, value=nome)
        ultima = 1
    existentes = set()
    for dia, mes, ano in ws.iter_rows(min_row=2, max_row=ultima, max_col=3, values_only=True):
        chave = (cot._inteiro(dia), cot._inteiro(mes), cot._inteiro(ano))
        if None not in chave:
            existentes.add(chave)
    return wb, ws, ultima, existentes


def salvar(wb, ws, ultima):
    """Salva a planilha; se o Excel estiver com ela aberta, espera e tenta de novo."""
    ws.auto_filter.ref = f"A1:E{ultima}"
    for tentativa in range(1, 7):
        try:
            wb.save(cot.ARQUIVO)
            return
        except PermissionError:
            print(f"  ! Não consegui salvar: feche {cot.ARQUIVO.name} no Excel. "
                  f"Nova tentativa em 15s ({tentativa}/6)...")
            time.sleep(15)
    print("Não foi possível salvar. O que faltou será buscado de novo na próxima execução.")
    sys.exit(1)


def linhas_do_dia(tabela, d):
    """Transforma a tabela do site em linhas [dia, mês, ano, produto, M.C.]."""
    if "M.C." not in tabela.columns or "Produto" not in tabela.columns:
        raise ValueError(f"layout diferente do esperado (colunas: {list(tabela.columns)})")
    base = tabela[["Produto", "M.C."]].dropna(subset=["M.C."])
    if cot.PRODUTOS:
        base = base[base["Produto"].isin(cot.PRODUTOS)]
    return [[d.day, d.month, d.year, p, float(mc)]
            for p, mc in zip(base["Produto"], base["M.C."])]


# ------------------------------------------------------------- principal ----
def processar(consultar, datas, ws, ultima, existentes, wb, limite=None):
    """
    Laço principal. `consultar(valor)` devolve a tabela do dia (ou None).
    Devolve (gravados, sem_dados, falhas, ultima).
    """
    sem_dados = set(SEM_DADOS.read_text().split()) if SEM_DADOS.exists() else set()
    pendentes = [(d, v) for d, v in datas
                 if (d.day, d.month, d.year) not in existentes
                 and d.strftime("%d/%m/%Y") not in sem_dados]
    if limite:
        pendentes = pendentes[:limite]

    print(f"Datas no site: {len(datas)} | já na planilha: {len(existentes)} | "
          f"a buscar agora: {len(pendentes)}")
    if not pendentes:
        return 0, 0, [], ultima

    gravados, novos_sem_dados, falhas, desde_salvar = 0, 0, [], 0
    vazios_pendentes, problemas_seguidos = [], 0
    inicio = time.time()

    def confirmar_vazios():
        """Datas vazias só viram 'sem boletim' quando o site provou estar saudável depois."""
        nonlocal novos_sem_dados
        if vazios_pendentes:
            with SEM_DADOS.open("a") as f:
                f.write("\n".join(vazios_pendentes) + "\n")
            novos_sem_dados += len(vazios_pendentes)
            vazios_pendentes.clear()

    try:
        for i, (d, valor) in enumerate(pendentes, start=1):
            texto = d.strftime("%d/%m/%Y")
            tabela, erro = None, None
            for tentativa in range(1, TENTATIVAS + 1):
                try:
                    tabela = consultar(valor)
                    erro = None
                    if tabela is not None:
                        break
                except Exception as e:  # noqa: BLE001 - qualquer falha de navegação
                    erro = e
                if tentativa < TENTATIVAS:
                    time.sleep(2 * tentativa)

            linhas = []
            if tabela is None:
                problemas_seguidos += 1
                if erro is not None:
                    falhas.append((texto, str(erro).splitlines()[0][:90]))
                    print(f"[{i}/{len(pendentes)}] {texto}: FALHOU ({falhas[-1][1]})")
                else:
                    vazios_pendentes.append(texto)
                    print(f"[{i}/{len(pendentes)}] {texto}: página sem tabela de produtos")
            else:
                try:
                    linhas = linhas_do_dia(tabela, d)
                except ValueError as e:
                    problemas_seguidos += 1
                    falhas.append((texto, str(e)[:90]))
                    print(f"[{i}/{len(pendentes)}] {texto}: FALHOU ({e})")
                    linhas = None
                if linhas is not None and not linhas:
                    print(f"[{i}/{len(pendentes)}] {texto}: nenhum produto para gravar")

            if problemas_seguidos >= MAX_PROBLEMAS_SEGUIDOS:
                print(f"\nPAREI: {problemas_seguidos} datas seguidas com problema. O site pode estar "
                      "fora do ar ou ter mudado de formato. Nada foi marcado como 'sem boletim'.\n"
                      "Tente de novo mais tarde; o robô continua de onde parou.")
                vazios_pendentes.clear()
                break

            if not linhas:
                continue

            confirmar_vazios()  # o site respondeu bem, então as datas vazias anteriores são reais
            problemas_seguidos = 0
            for linha in linhas:
                ultima += 1
                cot._escrever_linha(ws, ultima, linha)
            existentes.add((d.day, d.month, d.year))
            gravados += 1
            desde_salvar += 1

            decorrido = time.time() - inicio
            resta = decorrido / i * (len(pendentes) - i)
            print(f"[{i}/{len(pendentes)}] {texto}: {len(linhas)} produtos "
                  f"| faltam ~{resta / 60:.0f} min")

            if desde_salvar >= SALVAR_A_CADA:
                salvar(wb, ws, ultima)
                desde_salvar = 0
                print(f"  (planilha salva: {ultima} linhas)")
            time.sleep(PAUSA)
        else:
            confirmar_vazios()  # terminou a lista normalmente: vazios do final são reais
    except KeyboardInterrupt:
        print("\nInterrompido pelo usuário. Salvando o que já foi coletado...")
    finally:
        if desde_salvar:
            salvar(wb, ws, ultima)
    return gravados, novos_sem_dados, falhas, ultima


def main():
    limite = int(argumento("--limite")) if argumento("--limite") else None
    wb, ws, ultima, existentes = abrir_planilha()

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=not DEBUG)
        page = browser.new_page()
        datas = listar_datas(page)
        print(f"O site oferece {len(datas)} datas, de "
              f"{datas[-1][0]:%d/%m/%Y} a {datas[0][0]:%d/%m/%Y}.")

        gravados, sem_dados, falhas, ultima = processar(
            lambda valor: consultar_dia(page, valor),
            datas, ws, ultima, existentes, wb, limite)
        browser.close()

    print("\n=== Resumo ===")
    print(f"Dias gravados agora: {gravados} | sem boletim no site: {sem_dados} | falhas: {len(falhas)}")
    print(f"Planilha: {cot.ARQUIVO} ({ultima} linhas)")
    if falhas:
        print("Datas que falharam (rode de novo para tentar outra vez):")
        for texto, motivo in falhas:
            print(f"  {texto}: {motivo}")


if __name__ == "__main__":
    main()
